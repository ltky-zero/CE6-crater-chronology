from __future__ import annotations

import copy
import hashlib
import re
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "阿波罗地区定年-3.docx"
OUT_DOCX = ROOT / "阿波罗地区定年-4.docx"
OUT_DIR = Path(__file__).resolve().parent
AUDIT_DOCX = ROOT / "拟合记录" / "Yue与Neukum定年标定点数据对应及来源核查.docx"
FIG6_PREFIX = ROOT / "图6_五个候选来源区CSFD_AMA组合"
FIG7_PREFIX = ROOT / "图7_两套标定数据拟合_含CE6本地比较"


CE6L = {
    "name": "CE6_L",
    "age": 2.936,
    "age_err": 0.026,
    "n1": 3.16e-3,
    "n1_err": 0.11102566e-3,
}
CE6_LOCAL = {"name": "CE6 local (comparison only)", "age": 2.807, "age_err": 0.003, "n1": 2.08e-3, "n1_err": 0.13e-3}

NEUKUM_POINTS = [
    ("Highland (terrae)", 4.350, 0.100, 3.60e-1, 1.10e-1),
    ("Nectaris Basin (Apollo 16)", 4.100, 0.100, 1.20e-1, 0.40e-1),
    ("Apennines (Apollo 15)", 3.910, 0.100, 3.50e-2, 0.50e-2),
    ("Descartes Formation (Apollo 16)", 3.900, 0.100, 3.40e-2, 0.70e-2),
    ("Fra Mauro Formation (Apollo 14)", 3.910, 0.100, 3.70e-2, 0.70e-2),
    ("Taurus-Littrow Mare (Apollo 17)", 3.700, 0.100, 1.00e-2, 0.30e-2),
    ("Mare Tranquillitatis old (Apollo 11)", 3.720, 0.100, 9.00e-3, 1.80e-3),
    ("Mare Tranquillitatis young (Apollo 11)", 3.530, 0.050, 6.40e-3, 2.00e-3),
    ("Mare Imbrium (Apollo 15)", 3.280, 0.100, 3.20e-3, 1.10e-3),
    ("Oceanus Procellarum (Apollo 12)", 3.180, 0.100, 3.60e-3, 1.10e-3),
    ("Mare Fecunditatis (Luna 16)", 3.400, 0.040, 3.30e-3, 1.00e-3),
    ("Mare Crisium (Luna 24)", 3.300, 0.100, 3.00e-3, 0.60e-3),
    ("Copernicus (Apollo 12)", 0.850, 0.200, 1.30e-3, 0.30e-3),
    ("Tycho (Apollo 17)", 0.109, 0.004, 9.00e-5, 1.80e-5),
    ("North Ray (Apollo 16)", 0.0500, 0.0014, 4.40e-5, 1.10e-5),
    ("Cone (Apollo 14)", 0.0260, 0.0008, 2.10e-5, 0.50e-5),
    ("Phanerozoic craters, lunar equivalent", 0.375, 0.075, 3.60e-4, 1.10e-4),
]

YUE_POINTS = [
    ("Apollo 14", 3.922, 0.012, 3.82e-2, 1.07e-2),
    ("Apollo 17", 3.752, 0.007, 1.06e-2, 0.21e-2),
    ("Apollo 11", 3.578, 0.009, 6.64e-3, 0.561e-3),
    ("Apollo 15", 3.281, 0.012, 2.23e-3, 0.12e-3),
    ("Apollo 12", 3.242, 0.013, 2.34e-3, 0.05e-3),
    ("Luna 16", 3.382, 0.014, 4.32e-3, 0.01e-3),
    ("Luna 24", 3.328, 0.021, 2.54e-3, 0.08e-3),
    ("Copernicus", 0.800, 0.015, 6.68e-4, 0.048e-4),
    ("Tycho", 0.109, 0.004, 7.12e-5, 0.063e-5),
    ("North Ray", 0.0503, 0.0008, 3.90e-5, 0.043e-5),
    ("Cone", 0.0260, 0.0008, 2.10e-5, 0.50e-5),
    ("Chang'e-5", 2.030, 0.004, 1.74e-3, 0.022e-3),
    ("Chang'e-6 local", 2.807, 0.003, 2.08e-3, 0.13e-3),
    ("SPA basin", 4.247, 0.005, 3.69e-1, 0.48e-1),
]

P_NEUKUM = np.array([5.44e-14, 6.930, 8.380e-4])
P_NEUKUM_CE6L = np.array([4.76344828982205e-13, 6.328618583035769, 8.819129378948374e-4])
P_YUE = np.array([3.885e-15, 7.595, 7.377e-4])
P_YUE_CE6L = np.array([5.578e-15, 7.495, 7.532e-4])


def chronology(coefficients: np.ndarray, age: np.ndarray | float) -> np.ndarray:
    age = np.asarray(age, dtype=float)
    return coefficients[0] * np.expm1(coefficients[1] * age) + coefficients[2] * age


def unpack(points):
    return tuple(np.asarray([row[idx] for row in points], dtype=float) for idx in range(1, 5))


def figure_contract() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "图件合同_第四稿.md").write_text(
        "核心结论：组合图完整呈现五个候选来源区的CSFD/AMA；年代函数图以相同曲线和标定点直观比较Neukum与Yue体系。\n"
        "图件类型：定量网格。\n"
        "面板：图6包含CE6_F、CE6_L、Apollo Center、Apollo East、Apollo West五个原始CSFD/AMA结果；图7依次为Neukum全范围、Neukum局部、Yue全范围、Yue局部。\n"
        "数据完整性：图6仅重排已有五幅结果图；图7不改变两套拟合输入或系数，CE6本地仅在Neukum局部面板显示为比较点，不进入拟合。\n"
        "输出：600 dpi PNG、矢量PDF和可编辑文字SVG。\n",
        encoding="utf-8",
    )


def draw_figure6() -> None:
    mpl.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
    })
    sources = [
        ROOT / "定年结果" / "LITAO_Apollo_CE6_F.png",
        ROOT / "定年结果" / "LITAO_Apollo_CE6_L.png",
        ROOT / "定年结果" / "LITAO_Apollo_Center.png",
        ROOT / "定年结果" / "LITAO_Apollo_East.png",
        ROOT / "定年结果" / "LITAO_Apollo_West.png",
    ]
    images = [np.asarray(Image.open(path).convert("RGB")) for path in sources]
    fig = plt.figure(figsize=(7.15, 4.85), facecolor="white")
    grid = fig.add_gridspec(2, 6, wspace=0.015, hspace=0.02)
    positions = [grid[0, 0:2], grid[0, 2:4], grid[0, 4:6], grid[1, 1:3], grid[1, 3:5]]
    for image, position in zip(images, positions):
        axis = fig.add_subplot(position)
        axis.imshow(image)
        axis.axis("off")
    for suffix, kwargs in ((".png", {"dpi": 600}), (".pdf", {}), (".svg", {})):
        fig.savefig(FIG6_PREFIX.with_suffix(suffix), bbox_inches="tight", pad_inches=0.01, facecolor="white", **kwargs)
    plt.close(fig)


def draw_figure7() -> None:
    mpl.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 8.0,
        "axes.spines.right": False,
        "axes.spines.top": False,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
    })
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 6.3), constrained_layout=True)
    full_ages = np.linspace(0, 4.45, 1000)
    zoom_ages = np.linspace(2.75, 3.45, 500)
    systems = [
        ("Neukum Legacy Calibration Points", NEUKUM_POINTS, P_NEUKUM, P_NEUKUM_CE6L),
        ("Yue et al. (2026) Updated Calibration Points", YUE_POINTS, P_YUE, P_YUE_CE6L),
    ]
    for row_idx, (title, points, published, refitted) in enumerate(systems):
        age, age_err, n1, n1_err = unpack(points)
        for col_idx, (x_values, x_limits, full_range) in enumerate([
            (full_ages, (4.45, 0.0), True),
            (zoom_ages, (3.45, 2.75), False),
        ]):
            ax = axes[row_idx, col_idx]
            mask = np.ones_like(age, dtype=bool) if full_range else ((age >= 2.75) & (age <= 3.45))
            ax.errorbar(
                age[mask], n1[mask], xerr=age_err[mask], yerr=n1_err[mask],
                fmt="o", ms=3.3, color="#4D4D4D", ecolor="#A8A8A8", elinewidth=0.65,
                capsize=1.4, label="Calibration points", zorder=3,
            )
            ax.errorbar(
                CE6L["age"], CE6L["n1"], xerr=CE6L["age_err"], fmt="*", ms=9,
                color="#C43C39", ecolor="#C43C39", capsize=2.0, label="CE6_L", zorder=5,
            )
            if row_idx == 0 and col_idx == 1:
                ax.errorbar(
                    CE6_LOCAL["age"], CE6_LOCAL["n1"], xerr=CE6_LOCAL["age_err"],
                    yerr=CE6_LOCAL["n1_err"], fmt="o", ms=5.0, mfc="white", mec="#5B4B8A",
                    mew=1.25, ecolor="#5B4B8A", elinewidth=0.8, capsize=1.8,
                    label="CE6 local (comparison only)", zorder=6,
                )
            ax.plot(x_values, chronology(published, x_values), color="#2468A2", lw=1.5, label="Published curve")
            ax.plot(x_values, chronology(refitted, x_values), color="#C43C39", lw=1.5, label="Refit with CE6_L")
            ax.set_xlim(*x_limits)
            if full_range:
                ax.set_yscale("log")
                ax.set_ylim(1.2e-5, 8e-1)
                ax.set_title(f"{title}: Full Age Range", fontsize=8.5)
            else:
                ax.set_ylim(1.5e-3, 5.0e-3)
                ax.set_title(f"{title}: 2.75–3.45 Ga", fontsize=8.5)
            ax.set_xlabel("Age (Ga; older to younger)")
            ax.set_ylabel(r"Cumulative crater density $N(1)$ (km$^{-2}$)")
            ax.grid(alpha=0.18, which="both")
            ax.legend(loc="best", fontsize=6.45, frameon=False)
    for suffix, kwargs in ((".png", {"dpi": 600}), (".pdf", {}), (".svg", {})):
        fig.savefig(FIG7_PREFIX.with_suffix(suffix), bbox_inches="tight", facecolor="white", **kwargs)
    plt.close(fig)


def clear_paragraph_keep_properties(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def replace_paragraph(paragraph, text: str) -> None:
    original_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        original_rpr = copy.deepcopy(paragraph.runs[0]._r.rPr)
    clear_paragraph_keep_properties(paragraph)
    run = paragraph.add_run(text)
    if original_rpr is not None:
        run._r.insert(0, original_rpr)


def insert_paragraph_after(element, text: str, style_name: str | None = None):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    paragraph = Document().add_paragraph()  # a detached prototype for the wrapper type
    paragraph._p = new_p
    if style_name:
        ppr = new_p.get_or_add_pPr()
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), style_name)
        ppr.append(pstyle)
    paragraph.add_run(text)
    return paragraph


def paragraph_with_text(document: Document, exact_start: str):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(exact_start):
            return paragraph
    raise KeyError(exact_start)


def previous_paragraph_element(paragraph):
    previous = paragraph._p.getprevious()
    if previous is None:
        raise RuntimeError("Expected an image paragraph before its caption.")
    return previous


def add_picture_to_existing_paragraph(paragraph, image_path: Path, width_inches: float) -> None:
    clear_paragraph_keep_properties(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def remove_element(element) -> None:
    parent = element.getparent()
    parent.remove(element)


def shift_citation_numbers(text: str) -> str:
    def bracket_repl(match: re.Match[str]) -> str:
        inside = match.group(1)
        if not re.search(r"\d", inside):
            return match.group(0)
        return "[" + re.sub(r"\d+", lambda num: str(int(num.group(0)) + 1) if int(num.group(0)) >= 6 else num.group(0), inside) + "]"
    return re.sub(r"\[([^\]]+)\]", bracket_repl, text)


def shift_all_citations(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text:
            replace_paragraph(paragraph, shift_citation_numbers(paragraph.text))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text:
                        replace_paragraph(paragraph, shift_citation_numbers(paragraph.text))


def insert_reference(document: Document) -> None:
    reference_start = "[7] Xu, L."
    reference = paragraph_with_text(document, reference_start)
    new_p = OxmlElement("w:p")
    reference._p.addprevious(new_p)
    p = document.paragraphs[0]
    p._p = new_p
    p.style = reference.style
    p.add_run("[6] Wu, S. et al. Petrogenesis of Very-Low-Ti Basalts Returned by Chang'e-6 From the Lunar Farside. Geophysical Research Letters 52, e2025GL117921 (2025). https://doi.org/10.1029/2025GL117921.")


def update_references(document: Document) -> None:
    shift_all_citations(document)
    insert_reference(document)


def add_note_after_table(document: Document, table, text: str) -> None:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    wrapper = document.paragraphs[0]
    wrapper._p = new_p
    source_p = paragraph_with_text(document, "在2.936 Ga处")
    wrapper.style = source_p.style
    run1 = wrapper.add_run("注：")
    run1.bold = True
    wrapper.add_run(text)


def update_manuscript() -> None:
    document = Document(BASE)
    update_references(document)

    intro_anchor = paragraph_with_text(document, "Qian等[2]对阿波罗盆地内")
    new_p = OxmlElement("w:p")
    intro_anchor._p.addprevious(new_p)
    new_wrapper = document.paragraphs[0]
    new_wrapper._p = new_p
    new_wrapper.style = intro_anchor.style
    new_wrapper.add_run(
        "在样品来源方面，Cui等[3]已根据极低钛组分与着陆区附近月海单元的差异指出，这类碎屑可能来自着陆区东侧玄武岩单元；Wu等[6]对约2.9 Ga极低钛玄武岩的Pb-Pb年龄、化学组成及遥感资料进行综合比较，进一步提出其最可能来自着陆区东侧的A3单元。两项研究均给出可能的物源方向，而非对碎屑来源的唯一判定。"
    )

    fig6_caption = paragraph_with_text(document, "图6 CE6_F区域")
    fig6_image = previous_paragraph_element(fig6_caption)
    image_p = document.paragraphs[0]
    image_p._p = fig6_image
    add_picture_to_existing_paragraph(image_p, FIG6_PREFIX.with_suffix(".png"), 6.35)
    replace_paragraph(
        fig6_caption,
        "图6 五个候选来源区的撞击坑尺寸—频率分布（CSFD）及绝对模式年龄（AMA）。（a）CE6_F；（b）CE6_L；（c）Apollo Center；（d）Apollo East；（e）Apollo West。",
    )
    for start in ["图7 CE6_L区域", "图8 Apollo Center区域", "图9 Apollo East区域", "图10 Apollo West区域"]:
        caption = paragraph_with_text(document, start)
        image_element = previous_paragraph_element(caption)
        remove_element(image_element)
        remove_element(caption._p)

    fig7_caption = paragraph_with_text(document, "图11 Neukum旧参考点")
    fig7_image = previous_paragraph_element(fig7_caption)
    image_p = document.paragraphs[0]
    image_p._p = fig7_image
    add_picture_to_existing_paragraph(image_p, FIG7_PREFIX.with_suffix(".png"), 6.35)
    replace_paragraph(
        fig7_caption,
        "图7 Neukum旧参考点与Yue等更新参考点下的年代函数拟合。（a）Neukum旧参考点和经典曲线的完整年龄范围；（b）Neukum旧参考点与CE6_L候选标定点的2.75–3.45 Ga局部比较，空心圆为仅用于比较、不参与Neukum拟合的CE6本地标定点；（c）Yue等更新参考点和发表曲线的完整年龄范围；（d）Yue等更新参考点与CE6_L候选标定点的2.75–3.45 Ga局部比较。红线为加入CE6_L后的拟合曲线，蓝线为相应发表曲线。所有年龄横坐标均按左老右新排列。参考点来源见表3和表4[3,9,16–24]。",
    )

    replace_paragraph(
        paragraph_with_text(document, "五个候选来源区的统计结果见表2"),
        "五个候选来源区的统计结果见表2和图6。CE6_L的N(1)=3.16×10⁻³ km⁻²，为五个区域中的最高值；按Neukum年代函数得到的AMA为3.28（+0.028/−0.033）Ga，比2.936±0.026 Ga样品年龄老约0.34 Ga。Apollo West、Apollo Center和Apollo East的AMA均约为2.0 Ga或更年轻。",
    )
    replace_paragraph(
        paragraph_with_text(document, "按照2.3节所述方法"),
        "按照2.3节所述方法，本研究分别在上述两套标定数据中加入CE6_L候选标定点：Neukum分支由17个参考点增至18个，Yue分支由14个参考点增至15个。两套数据均以Neukum经典系数为初值，采用相同的未加权相对残差非线性最小二乘法。Yue等公开代码数组与论文表列的南极—艾特肯盆地年龄略有差别，本文采用论文表列的4.247 Ga，以保持Yue分支输入数据与论文表格一致。",
    )
    replace_paragraph(
        paragraph_with_text(document, "加入CE6_L候选标定点后，基于Neukum旧参考点"),
        "加入CE6_L候选标定点后，基于Neukum旧参考点的曲线在2.936 Ga处预测N(1)=2.65×10⁻³ km⁻²，相对残差RMS为18.1%，Shapiro–Wilk检验p=0.507，95%曲线区间为1.95–4.13×10⁻³ km⁻²；基于Yue更新参考点的曲线预测N(1)=2.23×10⁻³ km⁻²，相对残差RMS为15.8%，Shapiro–Wilk检验p=0.541，95%曲线区间为1.69–3.27×10⁻³ km⁻²。完整结果见图7和表4。",
    )
    replace_paragraph(
        paragraph_with_text(document, "CE6_L的AMA为3.28"),
        "CE6_L的AMA为3.28（+0.028/−0.033）Ga，比2.936±0.026 Ga样品年龄老约0.34 Ga。图7（b）和（d）及表3还显示，Apollo 12、Apollo 15、Luna 24和Luna 16等邻近标定点在2.9–3.4 Ga附近分布较为离散。CE6_L候选标定点补充了这一年龄段的样品年龄—撞击坑密度配对。",
    )
    replace_paragraph(
        paragraph_with_text(document, "采用相同的相对残差方法将CE6_L候选标定点"),
        "采用相同的相对残差方法将CE6_L候选标定点分别加入Neukum旧参考点和Yue更新参考点后，两套曲线在2.9–3.4 Ga附近均向较高N(1)方向调整，并保留了低年龄段线性项和高年龄段指数项的函数形式（图7）。",
    )
    replace_paragraph(
        paragraph_with_text(document, "在2.936 Ga处，加入CE6_L候选标定点后"),
        "在2.936 Ga处，加入CE6_L候选标定点后，Neukum分支的预测N(1)由2.50×10⁻³增至2.65×10⁻³ km⁻²，Yue分支由2.18×10⁻³增至2.23×10⁻³ km⁻²。图7显示了加入该点前后的曲线位置，表4列出了对应系数和预测值。",
    )

    heading42 = paragraph_with_text(document, "4.2 2.9–3.3 Ga年代关系")
    replace_paragraph(heading42, "4.2 2.8–3.4 Ga附近的年代关系")
    old42 = [
        paragraph_with_text(document, "Neukum旧参考点与Yue等[9]更新参考点"),
        paragraph_with_text(document, "两套重新拟合的曲线都将2.936 Ga极低钛"),
        paragraph_with_text(document, "综合化学组成、年龄和空间位置"),
        paragraph_with_text(document, "CE6_L的N(1)=3.16×10⁻³ km⁻²"),
        paragraph_with_text(document, "把CE6_L候选标定点分别加入Neukum旧参考点"),
    ]
    new_discussion = [
        "本研究以2.936±0.026 Ga样品年龄和CE6_L的N(1)=3.16×10⁻³ km⁻²组成候选标定点，并分别纳入Neukum旧参考点和Yue更新参考点。两套数据采用相同的函数形式和相对残差拟合方法，但标定点的年龄、N(1)和计数单元并不完全相同，因此得到的2.8–3.4 Ga附近年代关系也不相同（图7、表3）。",
        "在Neukum旧参考点组中，CE6_L与Apollo 12、Apollo 15、Luna 24和Luna 16的中心值集中在3.0–3.6×10⁻³ km⁻²，旧数据的误差范围大幅重叠，在图7（b）中形成近水平分布。作为比较点加入图7（b）的CE6本地玄武岩在2.807 Ga的N(1)=2.08×10⁻³ km⁻²，而CE6_L在2.936 Ga为3.16×10⁻³ km⁻²，中心值在0.129 Ga内增加约52%。在样品—地质单元配对成立的前提下，这一排列支持“2.9–3.4 Ga相对平坦、2.8–2.9 Ga快速转折”的候选解释；CE6本地点不参与Neukum拟合，仅用于展示这一局部对比。",
        "Yue等[9]在更新体系中采用了更精确的放射性年龄，并对计数区域进行重新选择。Apollo 12、Apollo 15和Luna 24的N(1)相对旧值降低，而Luna 16提高到4.32×10⁻³ km⁻²，因而不再呈现Neukum旧参考点的近水平关系。Yue等的补充材料表明，Luna 16的4.32×10⁻³ km⁻²来自光谱和地形较均一的计数区及LROC NAC制图；同一资料同时列出Werner等[21]得到的2.46±0.09×10⁻³ km⁻²，后者因计数面积较小、数据点偏离等时线而未被采用。Luna 16样品还记录了多期玄武岩活动，这使样品年龄与表面计数区的配对仍需结合地质背景持续检验[20,22]。",
        "因此，Yue体系的更新逻辑和大部分数据来源具有明确依据，不宜据此整体否定；但在2.8–3.4 Ga附近，Luna 16的计数区选择会直接影响局部曲线形态。若Yue更新点及其样品—单元对应关系均成立，则该区间可能经历多次撞击通量变化，单一平滑单调函数难以通过全部中心值；若进一步考虑替代计数结果、外来溅射物质、多期玄武岩和区域划分等系统差异，则目前尚不能唯一确定变化的具体形式。CE6_L候选标定点改变了两套标定数据在2.9–3.4 Ga附近的排列关系，后续仍需要具有明确物源的独立样品年龄和撞击坑统计结果加以约束。",
    ]
    for paragraph, text in zip(old42[:4], new_discussion):
        replace_paragraph(paragraph, text)
    remove_element(old42[4]._p)

    table3 = document.tables[2]
    add_note_after_table(
        document,
        table3,
        "Neukum旧参考点组与Yue更新参考点组的逐点对应、更新原因和来源核查见“Yue与Neukum定年标定点数据对应及来源核查”。CE6本地玄武岩仅作为图7（b）的比较点，不参与Neukum分支拟合。",
    )
    document.save(OUT_DOCX)


def remove_document_body(document: Document) -> None:
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_cell_text(cell, text: str, bold=False, font_size=8.0, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, hex_fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_width(cell, cm: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(document: Document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, font_size=8.0, color=(255, 255, 255))
        shade_cell(table.rows[0].cells[idx], "4F81BD")
        set_cell_width(table.rows[0].cells[idx], widths[idx])
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, font_size=7.4)
            set_cell_width(cells[idx], widths[idx])
            if len(table.rows) % 2 == 1:
                shade_cell(cells[idx], "EDF3F8")
    return table


def audit_mapping_rows():
    return [
        ("Highland", "4.35±0.10", "(3.6±1.1)×10⁻¹", "删除", "—", "年龄和N(1)均存较大争议；后续重标定普遍不采用。"),
        ("Nectaris Basin (A16)", "4.10±0.10", "(1.2±0.4)×10⁻¹", "删除", "—", "Apollo 16样品与Nectaris/Descartes/Imbrium物质对应不唯一，年龄争议大。"),
        ("Apennines (A15)", "3.91±0.10", "(3.5±0.5)×10⁻²", "删除", "—", "Imbrium年龄及样品来源仍有争议，未作为更新控制点。"),
        ("Descartes Formation (A16)", "3.90±0.10", "(3.4±0.7)×10⁻²", "删除", "—", "地层与Nectaris、Imbrium抛射物的关系不清，无法给出唯一标定年龄。"),
        ("Fra Mauro (A14)", "3.91±0.10", "(3.7±0.7)×10⁻²", "保留并更新", "3.922±0.012；(3.82±1.07)×10⁻²", "采用较新的U–Pb年龄和较均一计数区；N(1)仍反映计数区选择的不确定性。"),
        ("Taurus-Littrow (A17)", "3.70±0.10", "(1.0±0.3)×10⁻²", "保留并更新", "3.752±0.007；(1.06±0.21)×10⁻²", "采用Pb–Pb年龄；样品群和地质单元对应较清楚。"),
        ("Apollo 11 old", "3.72±0.10", "(9.0±1.8)×10⁻³", "删除", "—", "旧样品群的年龄分布和地质对应不支持作为独立控制点。"),
        ("Apollo 11 young", "3.53±0.05", "(6.4±2.0)×10⁻³", "保留并更新", "3.578±0.009；(6.64±0.561)×10⁻³", "采用更新Pb–Pb年龄与相应计数结果。"),
        ("Apollo 15 mare", "3.28±0.10", "(3.2±1.1)×10⁻³", "保留并更新", "3.281±0.012；(2.23±0.12)×10⁻³", "年龄精度提高；采用与橄榄石标准玄武岩相匹配且避开前熔岩坑的计数区。"),
        ("Apollo 12 mare", "3.18±0.10", "(3.6±1.1)×10⁻³", "保留并更新", "3.242±0.013；(2.34±0.05)×10⁻³", "采用最新Pb–Pb年龄与成分均一、CSFD贴合等时线的计数区。"),
        ("Luna 16", "3.40±0.04", "(3.3±1.0)×10⁻³", "保留并更新", "3.382±0.014；(4.32±0.01)×10⁻³", "年龄为多期活动平均值；N(1)采用Yue 2025的同质区与NAC计数，局部曲线对此选择敏感。"),
        ("Luna 24", "3.30±0.10", "(3.0±0.6)×10⁻³", "保留并更新", "3.328±0.021；(2.54±0.08)×10⁻³", "考虑更新的40K衰变常数；计数区按光谱匹配Luna 24样品的玄武岩单元。"),
        ("Copernicus", "0.85±0.20", "(1.3±0.3)×10⁻³", "保留并更新", "0.800±0.015；(6.68±0.048)×10⁻⁴", "更新Ar–Ar年龄和精细计数结果。"),
        ("Tycho", "0.109±0.004", "(9.0±1.8)×10⁻⁵", "保留并更新", "0.109±0.004；(7.12±0.063)×10⁻⁵", "年龄不变，采用更新N(1)。"),
        ("North Ray", "0.0500±0.0014", "(4.4±1.1)×10⁻⁵", "保留并更新", "0.0503±0.0008；(3.90±0.043)×10⁻⁵", "年龄和N(1)均采用更精确的更新值。"),
        ("Cone", "0.0260±0.0008", "(2.1±0.5)×10⁻⁵", "保留", "0.0260±0.0008；(2.1±0.5)×10⁻⁵", "数值不变。"),
        ("Phanerozoic lunar equivalent", "0.375±0.075", "(3.6±1.1)×10⁻⁴", "删除", "—", "由地球显生宙坑密度换算，系统差异较大，不再作为月球直接标定点。"),
        ("—", "—", "—", "新增CE5", "2.030±0.004；(1.74±0.022)×10⁻³", "嫦娥五号Pb–Pb样品年龄与北部风暴洋计数区配对。"),
        ("—", "—", "—", "新增CE6本地", "2.807±0.003；(2.08±0.13)×10⁻³", "CE6本地玄武岩Pb–Pb年龄与阿波罗月海计数区配对。"),
        ("—", "—", "—", "新增SPA", "4.247±0.005；(3.69±0.48)×10⁻¹", "CE6苏长岩年龄与剔除月海后的SPA大盆地计数区配对。"),
    ]


def make_audit_report() -> None:
    document = Document(BASE)
    remove_document_body(document)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)

    title = document.add_paragraph()
    title.style = document.styles["Heading 1"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Yue与Neukum定年标定点数据对应及来源核查")
    title_run.bold = True
    document.add_paragraph("核查对象：Neukum旧标定体系、Yue等（2026）更新体系及其2.9–3.4 Ga附近数据。本文档用于追溯输入数据、解释更新原因，并界定它们对局部年代关系讨论的约束。")

    h = document.add_paragraph("1. 核查结论")
    h.style = document.styles["Heading 2"]
    document.add_paragraph(
        "Yue等（2026）沿用Neukum的指数项加线性项形式和Neukum生产函数，但没有简单替换系数，而是重建了标定数据集：17个旧点中，11个被保留并更新，6个因年龄或样品—地质单元对应争议被删除；同时加入CE5、CE6本地玄武岩和SPA三个直接的返回样品约束。更新主要来自较精确的同位素年龄、与样品成分相匹配的地质单元，以及重新判读的撞击坑统计区。"
    )
    document.add_paragraph(
        "该体系的更新理由总体明确，但2.9–3.4 Ga局部形态对Luna 16的N(1)选择最为敏感。Yue等采用4.32±0.01×10⁻³ km⁻²；同一补充材料同时列出Werner等（2023）的2.46±0.09×10⁻³ km⁻²。前者来自更大、光谱和地形更均一的计数区，后者因统计面积较小且数据点偏离等时线而未被采用。因而，Yue更新体系可以作为一套有依据的标定数据，但不能单独用来唯一确定2.8–3.4 Ga区间是否存在平坦段、快速转折或多次变化。"
    )

    h = document.add_paragraph("2. 旧点与更新点的逐点对应")
    h.style = document.styles["Heading 2"]
    document.add_paragraph("N(1)单位均为km⁻²。旧值来自Yue等（2026）表1所列的Neukum体系；更新值、删除理由与计数区说明依据Yue等（2026）正文和补充材料。")
    add_table(
        document,
        ["Neukum旧点", "旧年龄（Ga）", "旧N(1)", "Yue处理", "Yue年龄；N(1)", "更新原因、年龄/计数区来源"],
        audit_mapping_rows(),
        [3.1, 2.0, 2.2, 2.1, 3.0, 12.0],
    )

    h = document.add_paragraph("3. 2.9–3.4 Ga附近的关键变化")
    h.style = document.styles["Heading 2"]
    key_rows = [
        ("Apollo 12", "3.18±0.10；3.60±1.10", "3.242±0.013；2.34±0.05", "年龄改用Snape等Pb–Pb结果；N(1)采用Werner等与样品成分匹配的均一计数区。"),
        ("Apollo 15", "3.28±0.10；3.20±1.10", "3.281±0.012；2.23±0.12", "中心年龄接近而精度提高；N(1)下调以排除前熔岩形成的大坑。"),
        ("Luna 24", "3.30±0.10；3.00±0.60", "3.328±0.021；2.54±0.08", "年龄考虑更新的40K衰变常数；N(1)采用光谱匹配的Luna 24周边单元。"),
        ("Luna 16", "3.40±0.04；3.30±1.00", "3.382±0.014；4.32±0.01", "年龄取两期喷发年龄平均；N(1)采用Yue 2025同质区。可选Werner值为2.46±0.09。"),
        ("CE6本地", "—", "2.807±0.003；2.08±0.13", "Yue新增的月背返回样品标定点。"),
        ("CE6_L", "—", "2.936±0.026；3.16（本研究）", "候选标定点；缺少N(1)误差，不虚构纵向误差。"),
    ]
    add_table(
        document,
        ["参考点", "Neukum旧值：年龄；N(1)（10⁻³ km⁻²）", "Yue/本研究值：年龄；N(1)（10⁻³ km⁻²）", "数据变化及其影响"],
        key_rows,
        [2.8, 4.7, 4.7, 13.2],
    )

    h = document.add_paragraph("4. 对2.8–3.4 Ga年代关系的含义")
    h.style = document.styles["Heading 2"]
    document.add_paragraph(
        "Neukum旧点中，CE6_L与Apollo 12、Apollo 15、Luna 24和Luna 16的N(1)中心值集中在约3.0–3.6×10⁻³ km⁻²，且旧点误差范围重叠较大。与CE6本地2.807 Ga、2.08×10⁻³ km⁻²相比，CE6_L的中心值在0.129 Ga内增加约52%，构成“相对平坦段加快速转折”的候选解释。该解释依赖CE6_L物源与CE6本地、旧Apollo/Luna标定点均可代表相应表面单元的前提。"
    )
    document.add_paragraph(
        "Yue更新后，Apollo 12、Apollo 15和Luna 24的N(1)降低而Luna 16升高；若这些数据及其样品—单元配对均成立，局部数据可支持多次撞击通量变化，单一平滑曲线难以通过所有中心值。另一方面，Luna 16的替代N(1)、多期玄武岩活动和计数面积差异表明，局部形态仍具有系统不确定性。因此，本研究在论文中不把任一种形态写成已被证明的历史，而是将CE6_L作为可检验的新候选标定点。"
    )

    h = document.add_paragraph("5. 核查来源、获取状态与文件校验")
    h.style = document.styles["Heading 2"]
    source_rows = [
        ("Yue et al., 2026, Sci. Adv.", "10.1126/sciadv.ady9265", "主文与补充材料", "主文9页；补充材料41页；PDF签名和题名已核对", "文献下载/完善稿引用文献/14_Yue_2026_lunar_chronology.pdf；14_Yue_2026_supplement.pdf"),
        ("Neukum et al., 2001, Space Sci. Rev.", "10.1023/A:1011989004263", "主文", "32页；PDF题名和DOI已核对", "文献下载/完善稿引用文献/15_Neukum_2001_cratering_records.pdf"),
        ("Werner et al., 2023, PSJ Part 2", "10.3847/PSJ/acdc16", "数据在Yue补充材料中交叉核对", "Yue补充材料明确列出2.46±0.09和采用4.32±0.01的理由", "主文需补充下载；当前核查依据Yue补充材料"),
        ("Yue et al., 2025, Icarus", "10.1016/j.icarus.2024.116348", "数据在Yue补充材料中交叉核对", "作为Luna 16的4.32±0.01来源；需补充下载主文", "当前核查依据Yue补充材料"),
        ("Wu et al., 2025, GRL", "10.1029/2025GL117921", "补充表已下载；主文获取受限", "补充表MD5=53503d3df8da1527ce173da6117238c2；出版商主文下载返回403", "文献下载/完善稿引用文献/06_Wu_2025_Supplementary_Tables.xlsx"),
    ]
    add_table(document, ["文献", "DOI", "获取范围", "校验结果", "文件路径或失败说明"], source_rows, [4.5, 4.5, 3.0, 7.5, 7.9])

    h = document.add_paragraph("参考文献")
    h.style = document.styles["Heading 2"]
    refs = [
        "Yue, Z. et al. Lunar chronology model with the Chang’e-6 farside samples and implications for the early impact history. Science Advances 12, eady9265 (2026). https://doi.org/10.1126/sciadv.ady9265.",
        "Neukum, G., Ivanov, B. A. & Hartmann, W. K. Cratering records in the inner solar system in relation to the lunar reference system. Space Science Reviews 96, 55–86 (2001). https://doi.org/10.1023/A:1011989004263.",
        "Werner, S. C., Bultel, B. & Rolf, T. Review and Revision of the Lunar Cratering Chronology—Lunar Timescale Part 2. The Planetary Science Journal 4, 147 (2023). https://doi.org/10.3847/PSJ/acdc16.",
        "Yue, Z. et al. New insights into the geological evolution history of Mare Fecunditatis. Icarus 425, 116348 (2025). https://doi.org/10.1016/j.icarus.2024.116348.",
        "Wu, S. et al. Petrogenesis of Very-Low-Ti Basalts Returned by Chang'e-6 From the Lunar Farside. Geophysical Research Letters 52, e2025GL117921 (2025). https://doi.org/10.1029/2025GL117921.",
        "Cui, Z. et al. A sample of the Moon's far side retrieved by Chang'e-6 contains 2.83-billion-year-old basalt. Science 386, 1395–1399 (2024). https://doi.org/10.1126/science.adt1093.",
        "Wang, C. et al. The source and thermal driver of young (<3.0 Ga) lunar volcanism. Science Advances 11, eadv9085 (2025). https://doi.org/10.1126/sciadv.adv9085.",
        "Snape, J. F. et al. Lunar basalt chronology, mantle differentiation and implications for determining the age of the Moon. Earth and Planetary Science Letters 451, 149–158 (2016). https://doi.org/10.1016/j.epsl.2016.07.026.",
        "Snape, J. F. et al. The timing of basaltic volcanism at the Apollo landing sites. Geochimica et Cosmochimica Acta 266, 29–53 (2019). https://doi.org/10.1016/j.gca.2019.07.042.",
        "Cohen, B. A. et al. Argon-40-argon-39 chronology and petrogenesis along the eastern limb of the Moon from Luna 16, 20 and 24 samples. Meteoritics & Planetary Science 36, 1345–1366 (2001). https://doi.org/10.1111/j.1945-5100.2001.tb01829.x.",
    ]
    for item in refs:
        p = document.add_paragraph(item)
        p.paragraph_format.space_after = Pt(3)
    AUDIT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(AUDIT_DOCX)


def check_output() -> None:
    for path in [BASE, OUT_DOCX, AUDIT_DOCX, FIG6_PREFIX.with_suffix(".png"), FIG7_PREFIX.with_suffix(".png")]:
        if not path.exists() or path.stat().st_size == 0:
            raise RuntimeError(f"Missing output: {path}")
    base_hash = hashlib.sha256(BASE.read_bytes()).hexdigest().upper()
    print("BASE_SHA256", base_hash)
    print("OUTPUT", OUT_DOCX)
    print("AUDIT", AUDIT_DOCX)


def main() -> None:
    figure_contract()
    draw_figure6()
    draw_figure7()
    update_manuscript()
    make_audit_report()
    check_output()


if __name__ == "__main__":
    main()
